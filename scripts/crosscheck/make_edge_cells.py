"""Writes tests/fixtures/crosscheck/edge-cells-openpyxl.xlsx with openpyxl, an independent .xlsx writer.

Each row of the "Cells" sheet holds one awkward kind of cell in column B, described in column A.
SimuLab's reader tests check that every kind is recognised. Requires: pip install openpyxl
"""
import datetime
import sys

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "Cells"
rows = [
    ("plain number", 8),
    ("percent format", 0.08),
    ("quoted percent sign", 7.5),
    ("date", datetime.datetime(2026, 1, 1)),
    ("boolean", True),
    ("error value", "#DIV/0!"),
    ("formula, no saved result", "=B2*2"),
    ("text with spaces", "  spaced  "),
    ("rich text", CellRichText([TextBlock(InlineFont(b=True), "Bold"), " and plain"])),
    ("ampersand and angle brackets", "A & B <c>"),
]
for index, (label, value) in enumerate(rows, start=1):
    ws.cell(row=index, column=1, value=label)
    ws.cell(row=index, column=2, value=value)
ws["B2"].number_format = "0%"
ws["B3"].number_format = '0.0"%"'
ws["B4"].number_format = "dd/mm/yyyy"
wb.save(out)
print("wrote", out)
